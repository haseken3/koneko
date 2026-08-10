"""PPTX のノート欄を読む共通ロジック（ナレーションカウンター / IDチェッカー 共用）。

2026-08-11 実事象（NeKo側で実測）: 教員が PowerPoint でノート枠を消したスライドは
`has_notes_slide` が True を返すのに `notes_text_frame` が None になる。
`slide.notes_slide.notes_text_frame.text` と書くとそこで AttributeError で落ちる
（実測: 品質管理2_2ブロック260810・全43枚の5枚目。教員提出73件2124スライド中1枚）。

読み手が2箇所（narration_counter / idcheck.pptx_reader）あるので、判断はこの1本に集約する
（別々に書くと片方だけ直る）。
"""

from pptx.enum.shapes import PP_PLACEHOLDER

# ノート面に残っていても原稿ではないプレースホルダ（スライド画像・ページ番号）。
# ⚠ ここを広げると原稿の欠落を黙って飲み込む。実測では本文枠を失った枚のノート面は
# ['', '5'] で、素朴に全テキストを連結すると "5" が原稿として数えられた。
_NON_NARRATION_NOTES_PLACEHOLDERS = (
    PP_PLACEHOLDER.SLIDE_IMAGE,
    PP_PLACEHOLDER.SLIDE_NUMBER,
)


class NotesBodyOrphanError(Exception):
    """ノート本文枠が無いのに、ノート面に本文らしいテキストが残っている。

    黙って空文字にすると尺が過小になり、IDチェッカーは「ナレーションなし」という
    事実と違う所見を教員に返す。数え落としを作らないため、ここは止める。
    """


def _orphan_notes_texts(notes_slide) -> list:
    """本文枠を失ったノート面に残っている「本文らしいテキスト」を集める。"""
    texts = []
    for shape in notes_slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type in _NON_NARRATION_NOTES_PLACEHOLDERS:
            continue
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if text:
            texts.append(text)
    return texts


def read_notes(slide, *, slide_num: int, total_slides: int) -> str:
    """1枚のノート本文を返す。ノート面／本文枠が無ければ空文字。

    本文枠が無い枚は「ナレーション無し」として通す（実測ではその枚に原稿が無い）。
    ただしノート面に本文らしいテキストが残っている個体だけは
    NotesBodyOrphanError で止める（数え落とし＝静かに間違った尺・所見を作らないため）。
    """
    try:
        if not slide.has_notes_slide:
            return ""
        notes_slide = slide.notes_slide
        frame = notes_slide.notes_text_frame
        if frame is not None:
            return frame.text.strip()
        orphans = _orphan_notes_texts(notes_slide)
    except Exception as e:
        raise RuntimeError(
            f"{total_slides}枚中の{slide_num}枚目でノート欄の読み取りに失敗しました: "
            f"{type(e).__name__}: {e}"
        ) from e

    if orphans:
        raise NotesBodyOrphanError(
            f"{total_slides}枚中の{slide_num}枚目: ノート欄の本文枠が無いのに、"
            f"ノート面に文章が残っています（このままだと数え落とします）: {orphans}"
        )
    return ""
