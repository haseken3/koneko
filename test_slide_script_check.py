"""スライド↔原稿チェックの回帰テスト（合成PPTX＋モックLLM・API不要・$0）。

⚠ このテストが測っているのは**配管**です。「AIの検出力」は測っていません
（LLM 応答はモック）。プロンプトが制作室の手直しを再現できるかは、実サンプルでの
実走で確かめること。ここで固定するのは次の5点:

  1. 本文抽出が group shape の中と表のセルまで届き、画面の上→下の順に並ぶ
     — 実測（第6回Step1・7枚）では本文のほぼ全部が group の中にあり、
       1段しか見ない読み方では表紙の抽出件数が 0 だった。この土台が抜けると
       「原稿にあってスライドに無い」が全枚で誤検知になる
  2. 全枚に出る定型行（ヘッダー/フッター）を突合対象から外す
  3. LLM が返した所見のうち、**宛先が確かめられないもの**を教員に出さない
     （実在しない引用の誤字／原稿が空の枚の不一致）
  4. 整った資料に配管が所見を生やさない（正例＝教員の正しい資料を赤だらけにしない）
  5. 1チャンクが落ちても他の結果は返り、落ちたことは黙らず errors に出る

実行: /usr/bin/python3 -m pytest test_slide_script_check.py -q （koneko ルートから）
"""

import sys
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import slide_script_check as ssc

HEADER = "情報化社会と情報倫理｜第6回：プライバシーと個人情報保護"


def _add_textbox(container, text, *, top_in, left_in=1.0, width_in=10.0):
    box = container.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(0.8))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    return box


def _make_deck(spec, *, header=HEADER, with_group=True) -> str:
    """spec = [(title, [bullets], notes), ...] から PPTX を作る。

    タイトルは group shape の中に入れる（実物のテンプレートと同じ形）。
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    blank = prs.slide_layouts[6]
    for title, bullets, notes in spec:
        slide = prs.slides.add_slide(blank)
        _add_textbox(slide.shapes, header, top_in=0.2)
        if with_group:
            group = slide.shapes.add_group_shape()
            _add_textbox(group.shapes, title, top_in=1.0)
        else:
            _add_textbox(slide.shapes, title, top_in=1.0)
        for i, bullet in enumerate(bullets):
            _add_textbox(slide.shapes, bullet, top_in=2.4 + i * 1.1)
        slide.notes_slide.notes_text_frame.text = notes
    path = Path(tempfile.mkdtemp()) / "deck.pptx"
    prs.save(str(path))
    return str(path)


# 整合デッキ（正例）: スライドの箇条書きが原稿の要点をすべて拾っている
ALIGNED_SPEC = [
    ("プライバシーとは何か",
     ["プライバシーとは秘密ではなく、自分の情報を自分でコントロールする権利である",
      "検索履歴・位置情報・購買データなど無形のデータもプライバシーに含まれる"],
     "プライバシーとは秘密のことではありません。自分の情報を自分でコントロールする権利です。"
     "検索履歴や位置情報、購買データのような無形のデータもプライバシーに含まれます。"),
    ("個人情報保護とは何か",
     ["個人情報保護はプライバシーという権利を守る社会の仕組みである",
      "法的規制だけでなく組織的運用と説明責任が不可欠である"],
     "個人情報保護は、プライバシーという権利を守るための社会の仕組みです。"
     "法的規制だけでなく、組織的な運用と説明責任が不可欠です。"),
    ("まとめ",
     ["技術進歩と人権保護の両輪で持続可能な社会が実現する"],
     "技術の進歩と人権の保護、この両輪が回って初めて持続可能な情報社会が実現します。"),
]

# 手直し前デッキ（負例）: S2 の原稿の結論がスライドに無い／S1 の本文に誤変換がある。
# S3 は整っている（所見が出る枚と出ない枚が混ざったデッキ＝実物に近い形）。
GAPPY_SPEC = [
    ("プライバシーとは何か",
     ["プライバシーとは秘密ではなく、自分の情報を自分でコントロールする権利である",
      "個人情報の保障は組織の責務である"],
     "プライバシーとは秘密のことではありません。自分の情報を自分でコントロールする権利です。"
     "個人情報の保障は組織の責務です。"),
    ("個人情報保護とは何か",
     ["個人情報保護はプライバシーという権利を守る社会の仕組みである"],
     "個人情報保護は、プライバシーという権利を守るための社会の仕組みです。"
     "説明責任を果たせない組織は存続できず、個人情報保護は組織の生存戦略です。"),
    ("まとめ",
     ["技術進歩と人権保護の両輪で持続可能な社会が実現する"],
     "技術の進歩と人権の保護、この両輪が回って初めて持続可能な情報社会が実現します。"),
]

_MISSING_QUOTE = "説明責任を果たせない組織は存続できず、個人情報保護は組織の生存戦略です。"
_TYPO_QUOTE = "個人情報の保障は組織の責務である"


def _finding(**kw):
    base = {
        "slide_num": 1, "kind": "不足", "where": "notes", "quote": "", "counterpart": "",
        "reason": "原稿にある要点がスライドに見当たりません。",
        "suggestion": "この一文をスライドに1行足すと揃います。", "confidence": "高",
    }
    base.update(kw)
    return base


def _mock_llm(consistency=None, typos=None, fail_on=None, record=None):
    """run_checks に差し込む LLM スタブ。tool 名でパスを振り分ける。"""
    def call(*, api_key, model, system, tool, user_prompt, max_tokens):
        if record is not None:
            record.append({"tool": tool["name"], "model": model, "prompt": user_prompt})
        if fail_on and fail_on in user_prompt:
            raise RuntimeError("APIが応答しませんでした")
        if tool["name"] == ssc.CONSISTENCY_TOOL["name"]:
            picked = [f for f in (consistency or []) if f"【S{f['slide_num']}】" in user_prompt]
            by_slide = {}
            for f in picked:
                by_slide.setdefault(f["slide_num"], []).append(f)
            return {"slides": [{"slide_num": n, "script_key_points": ["要点"], "findings": fs}
                               for n, fs in by_slide.items()]}
        picked = [f for f in (typos or []) if f"【S{f['slide_num']}】" in user_prompt]
        return {"findings": picked}
    return call


# ─────────────────────────────────────────────
# 1. 抽出（group / 表 / 読み順）
# ─────────────────────────────────────────────
def test_flat_reader_misses_group_text():
    """土台の確認: 1段しか見ない読み方では group 内の本文が取れない（実事象の再現）。"""
    slide = Presentation(_make_deck(ALIGNED_SPEC)).slides[0]
    flat = [sh.text_frame.text.strip() for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]
    assert "プライバシーとは何か" not in flat, "フィクスチャが group の入れ子を再現できていない"
    assert "プライバシーとは何か" in ssc.extract_slide_lines(slide)


def test_extract_reads_table_cells():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "目的限定"
    table.cell(0, 1).text = "必要な範囲だけ集める"
    lines = ssc.extract_slide_lines(slide)
    assert "目的限定 | 必要な範囲だけ集める" in lines


def test_group_child_coordinates_are_mapped_to_the_slide():
    """group の子座標（chOff）を変換せずに並べると読み順が壊れることを固定する。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(slide.shapes, "中段の箇条書き", top_in=3.0)
    group = slide.shapes.add_group_shape()
    _add_textbox(group.shapes, "画面のいちばん上の見出し", top_in=8.0)  # 子座標では下端
    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    xfrm.find(qn("a:off")).set("y", str(Inches(0.5)))  # 実際の配置は画面上端

    lines = ssc.extract_slide_lines(slide)
    assert lines == ["画面のいちばん上の見出し", "中段の箇条書き"]


def test_extract_slides_pairs_body_and_notes():
    slides = ssc.extract_slides(_make_deck(ALIGNED_SPEC))
    assert [s["slide_num"] for s in slides] == [1, 2, 3]
    assert slides[0]["title"] == "プライバシーとは何か", "定型ヘッダーがタイトルに化けている"
    assert "自分でコントロールする権利" in slides[0]["body"]
    assert slides[0]["notes"].startswith("プライバシーとは秘密")


# ─────────────────────────────────────────────
# 2. 定型行
# ─────────────────────────────────────────────
def test_boilerplate_is_detected_and_dropped_from_the_prompt():
    slides = ssc.extract_slides(_make_deck(ALIGNED_SPEC))
    boilerplate = ssc.detect_boilerplate(slides)
    assert HEADER in boilerplate
    assert "まとめ" not in boilerplate

    text = ssc.build_chunk_text(slides, boilerplate)
    assert HEADER not in text, "定型ヘッダーが突合対象に残っている（全枚が誤検知になる）"
    assert "プライバシーとは何か" in text
    assert "原稿（ノート欄）" in text


def test_short_deck_keeps_everything():
    """2枚デッキで「両方に出る行」を定型と決めつけない。"""
    slides = ssc.extract_slides(_make_deck(ALIGNED_SPEC[:2]))
    assert ssc.detect_boilerplate(slides) == []


# ─────────────────────────────────────────────
# 3. 正例（整った資料に所見を生やさない）
# ─────────────────────────────────────────────
def test_aligned_deck_produces_no_findings():
    result = ssc.run_checks(_make_deck(ALIGNED_SPEC), "dummy-key", call_llm=_mock_llm())
    assert result["consistency"] == []
    assert result["typos"] == []
    assert result["errors"] == []
    assert result["slides"] == 3


# ─────────────────────────────────────────────
# 4. 負例（仕込んだ不足・誤字が最後まで残る）
# ─────────────────────────────────────────────
def test_planted_gap_and_typo_survive_the_pipeline():
    result = ssc.run_checks(
        _make_deck(GAPPY_SPEC), "dummy-key",
        call_llm=_mock_llm(
            consistency=[_finding(slide_num=2, quote=_MISSING_QUOTE)],
            typos=[_finding(slide_num=1, kind="誤字脱字", where="slide",
                            quote=_TYPO_QUOTE, confidence="中",
                            suggestion="個人情報の保護は組織の責務である")],
        ),
    )
    assert len(result["consistency"]) == 1
    gap = result["consistency"][0]
    assert gap["slide_num"] == 2 and gap["kind"] == "不足"
    assert gap["quote_verified"] is True
    assert gap["slide_title"] == "個人情報保護とは何か"

    assert len(result["typos"]) == 1
    assert result["typos"][0]["quote"] == _TYPO_QUOTE


# ─────────────────────────────────────────────
# 5. 宛先が確かめられない所見を教員に出さない
# ─────────────────────────────────────────────
def test_typo_quote_that_is_nowhere_in_the_deck_is_dropped():
    result = ssc.run_checks(
        _make_deck(GAPPY_SPEC), "dummy-key",
        call_llm=_mock_llm(typos=[_finding(slide_num=1, kind="誤字脱字", where="slide",
                                           quote="この文はどこにも書かれていません")]),
    )
    assert result["typos"] == [], "場所を特定できない校正が教員に出ている"


def test_paraphrased_gap_quote_is_kept_but_flagged():
    """不足の引用は要約になりやすい。捨てずに「原稿の要点」として残す。"""
    result = ssc.run_checks(
        _make_deck(GAPPY_SPEC), "dummy-key",
        call_llm=_mock_llm(consistency=[
            _finding(slide_num=2, quote="説明責任を果たせない組織は存続できない（要約）"),
        ]),
    )
    assert len(result["consistency"]) == 1
    assert result["consistency"][0]["quote_verified"] is False


def test_slide_without_script_gets_no_consistency_finding():
    deck = _make_deck([("表紙", ["第6回 プライバシーと個人情報保護"], ""),
                       ("本編", ["自己情報コントロール権"], "自己情報コントロール権の話をします。")])
    result = ssc.run_checks(
        deck, "dummy-key",
        call_llm=_mock_llm(consistency=[
            _finding(slide_num=1, kind="過剰", where="slide",
                     quote="第6回 プライバシーと個人情報保護"),
        ]),
    )
    assert result["consistency"] == [], "原稿が無い枚に不一致所見が出ている"


def test_findings_are_capped_per_slide():
    quotes = ["プライバシーとは秘密のことではありません。",
              "自分の情報を自分でコントロールする権利です。",
              "個人情報の保障は組織の責務です。"]
    many = [_finding(slide_num=1, quote=q, confidence=c)
            for q, c in zip(quotes, ("低", "中", "高"))]
    many += [_finding(slide_num=1, quote="存在しない要点A", confidence="高"),
             _finding(slide_num=1, quote="存在しない要点B", confidence="高")]
    result = ssc.run_checks(_make_deck(GAPPY_SPEC), "dummy-key",
                            call_llm=_mock_llm(consistency=many))
    assert len(result["consistency"]) == ssc.MAX_FINDINGS_PER_SLIDE
    assert result["consistency"][0]["confidence"] == "高", "確度の高い順に残っていない"


# ─────────────────────────────────────────────
# 6. チャンク分割・並列・失敗の扱い
# ─────────────────────────────────────────────
def test_every_slide_is_sent_to_both_passes():
    spec = [(f"見出し{i}", [f"要点{i}"], f"原稿{i}です。") for i in range(1, 21)]
    record = []
    ssc.run_checks(_make_deck(spec), "dummy-key", chunk_size=8,
                   call_llm=_mock_llm(record=record))

    assert len(record) == 6, "20枚 / 8枚チャンク × 2パス = 6回になっていない"
    for tool_name in (ssc.CONSISTENCY_TOOL["name"], ssc.TYPO_TOOL["name"]):
        sent = {n for call in record if call["tool"] == tool_name
                for n in range(1, 21) if f"【S{n}】" in call["prompt"]}
        assert sent == set(range(1, 21)), f"{tool_name} に渡らなかった枚がある: {sorted(set(range(1, 21)) - sent)}"


def test_typo_pass_can_use_a_different_model():
    record = []
    ssc.run_checks(_make_deck(ALIGNED_SPEC), "dummy-key",
                   model="model-a", typo_model="model-b", call_llm=_mock_llm(record=record))
    used = {call["tool"]: call["model"] for call in record}
    assert used[ssc.CONSISTENCY_TOOL["name"]] == "model-a"
    assert used[ssc.TYPO_TOOL["name"]] == "model-b"


def test_each_finding_records_the_model_that_produced_it():
    """どのモデルで出た所見かを所見自身に刷る（モデル選定を実測で決めるための署名）。"""
    result = ssc.run_checks(
        _make_deck(GAPPY_SPEC), "dummy-key", model="model-a", typo_model="model-b",
        call_llm=_mock_llm(
            consistency=[_finding(slide_num=2, quote=_MISSING_QUOTE)],
            typos=[_finding(slide_num=1, kind="誤字脱字", where="slide", quote=_TYPO_QUOTE)],
        ),
    )
    assert result["consistency"][0]["model"] == "model-a"
    assert result["typos"][0]["model"] == "model-b"
    assert result["models"] == {"consistency": "model-a", "typos": "model-b"}


def test_failed_chunk_is_reported_and_the_rest_survives():
    spec = [(f"見出し{i}", [f"要点{i}"], f"原稿{i}です。") for i in range(1, 17)]
    result = ssc.run_checks(
        _make_deck(spec), "dummy-key", chunk_size=8,
        call_llm=_mock_llm(consistency=[_finding(slide_num=10, quote="原稿10です。")],
                           fail_on="【S1】"),
    )
    assert len(result["errors"]) == 2, "落ちたチャンクが黙って消えている"
    assert "S1〜S8" in result["errors"][0]
    assert [f["slide_num"] for f in result["consistency"]] == [10]


def test_progress_runs_from_zero_to_all():
    seen = []
    ssc.run_checks(_make_deck(ALIGNED_SPEC), "dummy-key",
                   call_llm=_mock_llm(), progress=lambda done, total: seen.append((done, total)))
    assert seen[0] == (0, 2)
    assert seen[-1] == (2, 2)


def test_empty_deck_returns_empty_result():
    prs = Presentation()
    path = Path(tempfile.mkdtemp()) / "empty.pptx"
    prs.save(str(path))
    result = ssc.run_checks(str(path), "dummy-key", call_llm=_mock_llm())
    assert result["slides"] == 0
    assert (result["consistency"], result["typos"], result["errors"], result["boilerplate"]) == ([], [], [], [])
    assert set(result["models"]) == {"consistency", "typos"}, "0枚でも使ったモデルは残す"


# ─────────────────────────────────────────────
# 7. ensure_result（自動実行・課金ループ防止の核。2026-09-17）
# ─────────────────────────────────────────────
def test_ensure_result_skips_runner_when_result_is_fresh():
    """file_id が一致する成功結果があれば runner を呼ばない（無駄な再課金をしない）。"""
    from slide_script_check_ui import ensure_result
    calls = []
    state = {"file": "a.pptx:100", "consistency": []}
    result = ensure_result(state, "a.pptx:100", lambda: calls.append(1) or {"consistency": []})
    assert calls == [], "既に新鮮な結果があるのに runner が呼ばれている"
    assert result is state


def test_ensure_result_does_not_retry_failed_without_force():
    """前回失敗した結果は、force無しでは runner を呼び直さない（自動rerunでの課金ループを防ぐ）。"""
    from slide_script_check_ui import ensure_result
    calls = []
    state = {"file": "a.pptx:100", "failed": "RuntimeError: boom"}
    result = ensure_result(state, "a.pptx:100", lambda: calls.append(1) or {"consistency": []})
    assert calls == [], "failed状態なのにforce無しでrunnerが呼ばれている"
    assert result is state


def test_ensure_result_runs_and_stores_failure_when_force():
    """force=True では実行し、失敗も file_id 付きで state に焼く（成功扱いで握りつぶさない）。"""
    from slide_script_check_ui import ensure_result

    def _boom():
        raise RuntimeError("boom")

    result = ensure_result({"file": "a.pptx:100", "failed": "old"}, "a.pptx:100", _boom, force=True)
    assert result["file"] == "a.pptx:100"
    assert "RuntimeError: boom" in result["failed"]


def test_ensure_result_runs_when_file_id_changes():
    """file_id が変わったら（別ファイルに差し替わったら）自動で runner を呼ぶ。"""
    from slide_script_check_ui import ensure_result
    calls = []
    state = {"file": "a.pptx:100", "consistency": []}
    result = ensure_result(state, "b.pptx:200",
                           lambda: calls.append(1) or {"consistency": ["new"]})
    assert calls == [1]
    assert result == {"file": "b.pptx:200", "consistency": ["new"]}


# ─────────────────────────────────────────────
# 8. is_total_failure（全滅と部分失敗を取り違えない。2026-09-17 verifier full W-1）
# ─────────────────────────────────────────────
def test_is_total_failure_true_when_all_chunks_failed():
    """errorsがあり所見が1件も無ければ全滅——「気づいた点なし」の緑表示を出してはいけない。"""
    result = {"errors": ["S1〜S8 の不一致チェック（claude-opus-5）が失敗しました: boom"],
              "consistency": [], "typos": []}
    assert ssc.is_total_failure(result) is True


def test_is_total_failure_false_when_some_findings_survive():
    """一部のチャンクが失敗しても、他の所見が残っていれば全滅ではない（部分失敗）。"""
    result = {"errors": ["S1〜S8 が失敗しました"], "consistency": [{"slide_num": 10}], "typos": []}
    assert ssc.is_total_failure(result) is False


def test_is_total_failure_false_when_no_errors():
    """errorsが無ければ、所見が0件でも「整っていた」だけで全滅ではない。"""
    result = {"errors": [], "consistency": [], "typos": []}
    assert ssc.is_total_failure(result) is False


# ─────────────────────────────────────────────
# 9. make_file_id / widget_suffix（file_id統一・widget key汚染防止。2026-09-17 verifier W-3/W-4）
# ─────────────────────────────────────────────
def test_make_file_id_uses_size_attribute_when_available():
    from slide_script_check_ui import make_file_id

    class _FakeUpload:
        name = "a.pptx"
        size = 12345
    assert make_file_id(_FakeUpload(), "a.pptx") == "a.pptx:12345"


def test_make_file_id_falls_back_to_path_getsize():
    from slide_script_check_ui import make_file_id
    path = _make_deck(ALIGNED_SPEC)
    file_id = make_file_id(path, "deck.pptx")
    assert file_id.startswith("deck.pptx:")
    assert file_id != "deck.pptx:None"


def test_make_file_id_falls_back_to_label_only_when_size_unavailable():
    from slide_script_check_ui import make_file_id
    assert make_file_id("/no/such/path.pptx", "label-only") == "label-only"


def test_widget_suffix_differs_between_file_ids():
    """2ファイル間で接尾辞が変われば、multiselectのkeyも変わり、旧選択が引き継がれない。"""
    from slide_script_check_ui import widget_suffix
    a = widget_suffix("a.pptx:100")
    b = widget_suffix("b.pptx:200")
    assert a != b
    assert widget_suffix("a.pptx:100") == a, "同じfile_idなら同じ接尾辞（無用な再描画を避ける）"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
