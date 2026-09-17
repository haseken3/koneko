"""ko-NeKo — スライド本文とノート原稿の突き合わせ（不一致／誤字脱字の「可能性」提示）。

教員が資料を提出する前に自分で回す道具。**断定しない**のが仕様の中心にある：
このモジュールが返すのは「所見（気づいた点）」であって判定ではない。最終判断は教員。

設計（2026-09-17）:
- 2パス（不一致／誤字脱字）× チャンク（既定8枚）を1つの ThreadPoolExecutor に混ぜて並列投入する。
  40枚デッキでも壁時計時間はチャンク1本分に近づく。1チャンクが落ちても他は返す（errors に積む）。
- 誤検知の抑制は3段構え:
  (1) 機械: 全枚に出る定型行（ヘッダー/フッター/Step表記）を突合対象から外す。
      「原稿にあってスライドに無い要点」を探す機能なので、定型行を要点扱いすると全枚が汚れる。
  (2) スキーマ: 1枚あたりの所見を MAX_FINDINGS_PER_SLIDE 件に縛る。原稿は口語で長く
      スライドは要約なので、「不足」は縛らないと1枚10件になる。
  (3) プロンプト: 先に原稿の要点（script_key_points）を挙げさせ、所見はその中からだけ選ばせる。
      スライドを見てから粗探しするのではなく、要点を決めてから突き合わせる順序を強制する。
- スライド本文の抽出は group shape を再帰的に辿り、group の子座標変換（chOff/chExt）を
  かけてから上→下に並べる。実測（第6回Step1・7枚）では本文のほぼ全部が group の中にあり、
  1段しか見ない `idcheck/pptx_reader.read_slides` では S1 の抽出件数が 0 だった。
- ノート欄の読み取り判断は `pptx_notes.read_notes` に集約されたまま（ここには書かない）。
"""

import math
import re
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import anthropic
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

# flat（app.py）とパッケージ経由（koneko.ui）の両方で読まれるので、自分の居るディレクトリを
# 末尾に足してから flat import する（narration_counter.py と同じ作法）。
_HERE = str(Path(__file__).resolve().parent)
if _HERE not in sys.path:
    sys.path.append(_HERE)
from llm_models import MODEL_MISMATCH, MODEL_TYPO  # noqa: E402
from pptx_notes import read_notes  # noqa: E402

MAX_RETRIES = 3
CHUNK_SIZE = 8
MAX_WORKERS = 6
MAX_FINDINGS_PER_SLIDE = 3
MAX_TOKENS_CONSISTENCY = 8000
MAX_TOKENS_TYPO = 4000

KINDS_CONSISTENCY = ("不足", "相違", "過剰")
KIND_TYPO = "誤字脱字"
CONFIDENCES = ("高", "中", "低")
_CONF_RANK = {c: i for i, c in enumerate(CONFIDENCES)}
_KIND_RANK = {k: i for i, k in enumerate(KINDS_CONSISTENCY + (KIND_TYPO,))}

# 全枚に出る行を定型（ヘッダー/フッター）とみなす閾値。2枚デッキで誤って消さないよう下限も置く。
BOILERPLATE_MIN_SLIDES = 3
BOILERPLATE_RATIO = 0.6


# ─────────────────────────────────────────────
# 抽出
# ─────────────────────────────────────────────
def _child_transform(group_shape):
    """group の子座標を親座標へ移す関数を返す（a:chOff/a:chExt の逆変換）。"""
    grp_pr = group_shape._element.find(qn("p:grpSpPr"))
    xfrm = grp_pr.find(qn("a:xfrm")) if grp_pr is not None else None
    if xfrm is None:
        return lambda x, y: (x, y)
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return lambda x, y: (x, y)
    ox, oy = int(off.get("x")), int(off.get("y"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    cox, coy = int(ch_off.get("x")), int(ch_off.get("y"))
    ccx, ccy = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
    sx = cx / ccx if ccx else 1.0
    sy = cy / ccy if ccy else 1.0
    return lambda x, y: (ox + (x - cox) * sx, oy + (y - coy) * sy)


def _walk_shapes(shapes, transform):
    """(y, x, text) を再帰的に集める。group は子座標変換を合成して潜る。"""
    found = []
    for shape in shapes:
        x, y = transform(shape.left or 0, shape.top or 0)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            inner = _child_transform(shape)
            found.extend(_walk_shapes(
                shape.shapes,
                lambda a, b, outer=transform, inner=inner: outer(*inner(a, b)),
            ))
            continue
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                found.append((y, x, text))
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(t for t in cells if t)
                if line:
                    found.append((y, x, line))
    return found


def extract_slide_lines(slide) -> list:
    """1枚のスライド本文を、画面の上→下（同じ高さなら左→右）の順のテキスト行で返す。

    group shape・表のセルも辿る（`idcheck/pptx_reader.py` は1段しか見ず、
    本文が group に入っているテンプレートでは何も取れない）。

    `idcheck/pptx_reader.read_slides` をgroup対応へ拡張する案も検討したが見送った
    （2026-09-17・inspector指摘）: idcheckはタイトル生成・読み順を問わない集計が主用途で
    要求仕様が違い、共用にすると両者の互換維持コストが増える。ノート抽出（`pptx_notes.read_notes`）
    のように判断が1つしかないものは共有するが、本文抽出は用途ごとに独立させる方針。
    """
    items = sorted(_walk_shapes(slide.shapes, lambda x, y: (x, y)))
    return [text for _, _, text in items]


def extract_slides(pptx_file) -> list:
    """PPTX から全スライドの本文行・ノート原稿を抽出する。

    Returns:
        list[dict]: slide_num / title / body_lines / body / notes
    """
    prs = Presentation(pptx_file)
    total = len(prs.slides)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = extract_slide_lines(slide)
        slides.append({
            "slide_num": i,
            "title": lines[0].split("\n")[0][:120] if lines else "",
            "body_lines": lines,
            "body": "\n".join(lines),
            "notes": read_notes(slide, slide_num=i, total_slides=total),
        })
    # タイトルは「定型行を除いた先頭行」が本来の見出し。定型が決まってから貼り直す。
    boilerplate = set(detect_boilerplate(slides))
    for s in slides:
        real = [l for l in s["body_lines"] if l not in boilerplate]
        s["title"] = real[0].split("\n")[0][:120] if real else ""
    return slides


def detect_boilerplate(slides) -> list:
    """多くの枚に繰り返し出る行（ヘッダー/フッター/Step表記）を返す。"""
    total = len(slides)
    if total < BOILERPLATE_MIN_SLIDES:
        return []
    counts = Counter(line for s in slides for line in set(s["body_lines"]))
    threshold = max(BOILERPLATE_MIN_SLIDES, math.ceil(total * BOILERPLATE_RATIO))
    return [line for line, n in counts.items() if n >= threshold]


# ─────────────────────────────────────────────
# tool-use スキーマ
# ─────────────────────────────────────────────
def _finding_properties(kinds, *, quote_desc, counterpart_desc, suggestion_desc):
    return {
        "slide_num": {
            "type": "integer",
            "description": "所見の対象スライド番号（提示された S番号をそのまま使う）。",
            "minimum": 1,
        },
        "kind": {"type": "string", "enum": list(kinds), "description": "所見の種類。"},
        "where": {
            "type": "string",
            "enum": ["slide", "notes"],
            "description": "quote がどちら側の原文か。slide=スライド本文 / notes=ノート原稿。",
        },
        "quote": {"type": "string", "description": quote_desc},
        "counterpart": {"type": "string", "description": counterpart_desc},
        "reason": {
            "type": "string",
            "description": "なぜ気づいたのかを1〜2文で。教員が5秒で判断できる具体性で書く。",
        },
        "suggestion": {"type": "string", "description": suggestion_desc},
        "confidence": {
            "type": "string",
            "enum": list(CONFIDENCES),
            "description": (
                "高=原稿の明確な主張・数値・結論がスライドに全く見当たらない、または事実が食い違う。"
                "中=要点だが図表で表現されている可能性や解釈の余地がある。"
                "低=表現の差に近く、指摘しなくても支障がない。"
            ),
        },
    }


CONSISTENCY_TOOL = {
    "name": "report_consistency",
    "description": (
        "オンデマンド授業のスライド本文とノート原稿（ナレーション）を突き合わせ、"
        "内容が食い違っている可能性のある箇所を報告する。"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "slides": {
                "type": "array",
                "description": "提示されたスライド1枚につき1要素。順番は提示順。",
                "items": {
                    "type": "object",
                    "properties": {
                        "slide_num": {"type": "integer", "minimum": 1},
                        "script_key_points": {
                            "type": "array",
                            "description": (
                                "そのスライドの原稿（ノート欄）だけを読んで、学習者が持ち帰るべき要点を"
                                "2〜5個、原稿の言葉で書き出す。スライド本文は見ずに書くつもりで抽出する。"
                                "原稿が空の場合は空配列。"
                            ),
                            "maxItems": 5,
                            "items": {"type": "string"},
                        },
                        "findings": {
                            "type": "array",
                            "description": (
                                f"そのスライドの所見。最大{MAX_FINDINGS_PER_SLIDE}件、重要な順。"
                                "気づいた点が無ければ空配列（0件が正常な状態）。"
                            ),
                            "maxItems": MAX_FINDINGS_PER_SLIDE,
                            "items": {
                                "type": "object",
                                "properties": _finding_properties(
                                    KINDS_CONSISTENCY,
                                    quote_desc=(
                                        "根拠となる原文を**そのまま**抜き出す（要約しない・言い換えない）。"
                                        "不足なら原稿側の該当文、過剰・相違ならスライド側の該当行。"
                                    ),
                                    counterpart_desc=(
                                        "相手側の対応箇所の原文。相違なら原稿側の該当文を必ず入れる。"
                                        "不足・過剰で対応箇所が無い場合は空文字。"
                                    ),
                                    suggestion_desc=(
                                        "教員が取れる対応を1文で（例:「この一文をスライドに1行足すと原稿と揃います」）。"
                                        "命令口調にせず、提案の形にする。"
                                    ),
                                ),
                                "required": ["slide_num", "kind", "where", "quote",
                                             "counterpart", "reason", "suggestion", "confidence"],
                            },
                        },
                    },
                    "required": ["slide_num", "script_key_points", "findings"],
                },
            },
        },
        "required": ["slides"],
    },
}

TYPO_TOOL = {
    "name": "report_typos",
    "description": "スライド本文とノート原稿から、誤字・脱字・誤変換の可能性がある箇所を報告する。",
    "input_schema": {
        "type": "object",
        "properties": {
            "findings": {
                "type": "array",
                "description": (
                    f"所見の配列。1枚あたり最大{MAX_FINDINGS_PER_SLIDE}件。"
                    "気づいた点が無ければ空配列（0件が正常な状態）。"
                ),
                "items": {
                    "type": "object",
                    "properties": _finding_properties(
                        (KIND_TYPO,),
                        quote_desc=(
                            "誤りを含む最小の一文を、原文のまま**1文字も変えずに**抜き出す"
                            "（修正後の文を書かない）。"
                        ),
                        counterpart_desc="使わない。空文字にする。",
                        suggestion_desc="修正案の文を1つ（quote と同じ範囲を、直した形で書く）。",
                    ),
                    "required": ["slide_num", "kind", "where", "quote",
                                 "counterpart", "reason", "suggestion", "confidence"],
                },
            },
        },
        "required": ["findings"],
    },
}


# ─────────────────────────────────────────────
# プロンプト
# ─────────────────────────────────────────────
CONSISTENCY_SYSTEM = f"""あなたは大学のオンデマンド授業教材を点検する編集者です。

教員が作った PowerPoint には、スライド本文（受講者が画面で読む）と
ノート欄の原稿（教員が読み上げるナレーション）の2つが入っています。
あなたの仕事は、この2つの**内容**が食い違っている可能性のある箇所を見つけることです。

【所見の3種類】
- 不足: 原稿で述べている**要点**が、スライドのどこにも書かれていない
- 相違: 両方にあるが、事実・数値・固有名詞・断定の強さ・範囲が食い違っている
- 過剰: スライドに書かれている要点を、原稿が一言も触れていない（読み上げられない）

【「要点」の定義（ここが一番大事）】
要点とは、学習者がそのスライドで持ち帰るべき **主張・結論・定義・分類・数値・手順・次回予告** です。
次のものは要点ではありません。挙げてはいけません:
- 原稿の言い換え、具体例の細部、言い直し、つなぎの言葉
- 「〜してください」「意識してください」のような呼びかけ・励まし
- 原稿の方が詳しいのは当たり前です。**詳細度の差そのものを不足と呼ばない**

【やってはいけないこと（誤検知を避ける）】
- 表現・語順・文体の違いだけを「相違」にしない。要旨が同じなら食い違いではない
- スライドの箇条書きが**別の言い方で同じことを言っている**なら、不足ではない
- 図・グラフ・イラストで表現されている可能性があるものは、確信度を下げる（テキストだけでは見えない）
- 原稿が空のスライド（表紙・区切り）には所見を出さない
- 全スライドに共通して出る定型のヘッダー・フッター・Step表記は、要点として扱わない
- 1枚あたり最大{MAX_FINDINGS_PER_SLIDE}件。多く挙げるほど良いのではありません。
  **所見0件は正常な状態です**。整っているスライドには何も出さないでください

【手順】
1. まずそのスライドの**原稿だけ**を読み、要点を2〜5個書き出す（script_key_points）
   原稿の**結び（最後の1〜2文）**には、その枚の結論・規範的な主張・次回予告が置かれていることが
   多いので、要点を挙げるときに必ず一度確認する。
   ただし「以上で終わります」「ご清聴ありがとうございました」のような締めの挨拶は要点ではない
2. 次にスライド本文を読み、1で挙げた要点が載っているかを1つずつ確かめる
3. 載っていない／食い違う要点だけを findings に書く。不足・相違の quote は
   1で挙げた要点に対応する**原文**をそのまま抜き出す

【言葉づかい】
この結果は教員にそのまま表示されます。断定せず、「可能性」の提示にとどめてください。
reason は「〜が見当たりません」「〜と食い違って見えます」、
suggestion は「〜すると揃います」のような提案の形で書きます。

必ず report_consistency ツールで出力してください。"""

TYPO_SYSTEM = f"""あなたは大学のオンデマンド授業教材を校正する校正者です。

スライド本文とノート原稿（ナレーション）から、**誤字・脱字・誤変換**の可能性がある箇所を
見つけてください。

【拾うもの】
- 同音異義語の誤変換（例: 「以外」と「意外」、「保証」と「保障」の取り違え）
- 脱字（助詞や文字の抜け）、衍字（余分な文字の重複）
- 明らかに壊れた文（主語と述語がつながっていない、途中で切れている）
- 括弧・引用符の対応が取れていない、句読点の重複
- 数字・単位・年号の明らかな誤り（例: 同じ資料内で数値が食い違う）

【拾わないもの（ここを守らないと教員の正しい資料が赤だらけになります）】
- 表記ゆれ単体（「サーバ／サーバー」「行う／行なう」など）。誤りではありません
- 専門用語・学術用語・固有名詞・製品名。知らないだけの可能性があるので原則指摘しない
- 話し言葉・体言止め・倒置・強調のための繰り返し。原稿はナレーションなので口語で正常です
- 全角半角の混在、スペースの有無
- 改行・箇条書き・見出し記号などの体裁やレイアウト（文の切れ目に見えても誤りではありません）
- 文体の好み（です・ます／だ・である）の統一
- 1枚あたり最大{MAX_FINDINGS_PER_SLIDE}件。**所見0件は正常な状態です**

【quote の書き方】
誤りを含む最小の一文を、原文のまま1文字も変えずに抜き出してください
（修正後の文を quote に書くと、教員がどこの話か分からなくなります）。
suggestion には同じ範囲を直した形を書きます。

【言葉づかい】
この結果は教員にそのまま表示されます。断定せず「〜の可能性があります」の形で書いてください。

必ず report_typos ツールで出力してください。"""


def build_chunk_text(slides, boilerplate=()) -> str:
    """チャンクのスライド群を、LLM に渡す番号付きテキストへ整形する。"""
    skip = set(boilerplate)
    parts = []
    for s in slides:
        lines = [l for l in s["body_lines"] if l not in skip]
        body = "\n".join(f"  ・{l}" for l in lines) if lines else "  （本文テキストなし）"
        notes = s["notes"].strip()
        parts.append(
            f"【S{s['slide_num']}】\n"
            f"スライド本文:\n{body}\n"
            f"原稿（ノート欄）:\n  {notes if notes else '（原稿なし）'}\n"
        )
    return "\n".join(parts)


def _chunk_prompt(slides, boilerplate, total_slides, task_line) -> str:
    nums = ", ".join(f"S{s['slide_num']}" for s in slides)
    head = f"授業デッキ全{total_slides}枚のうち、{nums} を点検してください。\n"
    if boilerplate:
        listed = "／".join(f"「{b}」" for b in sorted(boilerplate)[:5])
        head += (
            f"※ 次の行は全スライド共通の定型（ヘッダー・フッター）なので本文から除いてあります: {listed}\n"
        )
    return f"{head}\n---\n{build_chunk_text(slides, boilerplate)}---\n\n{task_line}"


# ─────────────────────────────────────────────
# LLM 呼び出し（差し替え可能）
# ─────────────────────────────────────────────
def call_anthropic(*, api_key, model, system, tool, user_prompt, max_tokens) -> dict:
    """tool_use の input を返す。中身が不正な回（tool_use 欠落）も含めてリトライする。

    SDK は 429/5xx を自動リトライする。このループはそれでも落ちた場合に加えて
    「API は成功したが tool_use が返らなかった」を拾い直すため意図的に被せている
    （step_segmenter.segment_steps と同型）。
    """
    client = anthropic.Anthropic(api_key=api_key)
    last_err = None
    for _ in range(MAX_RETRIES):
        try:
            resp = client.messages.create(
                model=model,
                max_tokens=max_tokens,
                system=system,
                tools=[tool],
                tool_choice={"type": "tool", "name": tool["name"]},
                messages=[{"role": "user", "content": user_prompt}],
            )
            block = next(
                (b for b in resp.content if b.type == "tool_use" and b.name == tool["name"]),
                None,
            )
            if block is None:
                last_err = f"{tool['name']} ツール呼び出しが返らなかった"
                continue
            return block.input
        except (anthropic.APIStatusError, anthropic.APIConnectionError) as e:
            last_err = f"{type(e).__name__}: {e}"
            continue
    raise RuntimeError(f"AIの応答を取得できませんでした（{MAX_RETRIES}回試行）: {last_err}")


# ─────────────────────────────────────────────
# 所見の正規化
# ─────────────────────────────────────────────
def _normalize(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def _quote_found(quote: str, target: str) -> bool:
    """引用が原文に実在するか（空白・改行の差は無視）。"""
    n_quote = _normalize(quote)
    return bool(n_quote) and n_quote in _normalize(target)


def _sanitize_findings(raw, *, slide_map, allowed_kinds) -> list:
    """LLM 出力の findings を、教員に出してよい形に正規化する。

    捨てる／格下げするもの:
    - 対象チャンク外・未知の kind・quote が空（宛先不明の所見）
    - 原稿が空の枚の不一致所見（ナレーション無しは文字数カウンター側で分かる）
    - 原文に見つからない引用の誤字所見（場所が特定できない校正は使えない）
    不足・相違の引用は要約になりやすいので、捨てずに quote_verified=False を立てて残す。
    """
    cleaned = []
    for f in raw or []:
        if not isinstance(f, dict):
            continue
        num = f.get("slide_num")
        kind = (f.get("kind") or "").strip()
        if not isinstance(num, int) or num not in slide_map or kind not in allowed_kinds:
            continue
        slide = slide_map[num]
        if kind != KIND_TYPO and not slide["notes"].strip():
            continue
        quote = (f.get("quote") or "").strip()
        if not quote:
            continue
        where = (f.get("where") or "").strip()
        if where not in ("slide", "notes"):
            where = "notes" if kind == "不足" else "slide"
        verified = _quote_found(quote, slide["notes"] if where == "notes" else slide["body"])
        if kind == KIND_TYPO and not verified:
            continue
        confidence = (f.get("confidence") or "").strip()
        if confidence not in CONFIDENCES:
            confidence = "中"
        cleaned.append({
            "slide_num": num,
            "kind": kind,
            "where": where,
            "quote": quote,
            "counterpart": (f.get("counterpart") or "").strip(),
            "reason": (f.get("reason") or "").strip(),
            "suggestion": (f.get("suggestion") or "").strip(),
            "confidence": confidence,
            "quote_verified": verified,
            "slide_title": slide.get("title", ""),
            "model": f.get("model", ""),
        })

    cleaned.sort(key=lambda f: (f["slide_num"], _CONF_RANK[f["confidence"]],
                               _KIND_RANK.get(f["kind"], 9)))
    per_slide = Counter()
    capped = []
    for f in cleaned:
        if per_slide[f["slide_num"]] >= MAX_FINDINGS_PER_SLIDE:
            continue
        per_slide[f["slide_num"]] += 1
        capped.append(f)
    return capped


def _flatten_consistency(data) -> list:
    raw = []
    for entry in (data or {}).get("slides") or []:
        if not isinstance(entry, dict):
            continue
        for f in entry.get("findings") or []:
            if isinstance(f, dict):
                f.setdefault("slide_num", entry.get("slide_num"))
                raw.append(f)
    return raw


def is_total_failure(result) -> bool:
    """全チャンクが失敗し、所見が1件も無い状態か判定する（純関数）。

    errors が非空なのに consistency/typos が両方空だと、画面は「気づいた点は
    ありませんでした」という緑の安心表示を出してしまう——実際は「チェックできて
    いない」だけ（2026-09-17 verifier full 指摘）。呼び出し側はこれが True の時、
    成功として扱わず失敗経路（赤表示＋もう一度ボタン）に乗せること。
    """
    return bool(result.get("errors")) and not result.get("consistency") and not result.get("typos")


# ─────────────────────────────────────────────
# 実行
# ─────────────────────────────────────────────
def _chunks(slides, chunk_size):
    return [slides[i:i + chunk_size] for i in range(0, len(slides), chunk_size)]


def run_checks(pptx_file=None, api_key=None, *, slides=None, model=MODEL_MISMATCH,
               typo_model=MODEL_TYPO, chunk_size=CHUNK_SIZE, max_workers=MAX_WORKERS,
               call_llm=None, progress=None) -> dict:
    """不一致パスと誤字脱字パスを並列で回し、所見をまとめて返す。

    Args:
        pptx_file: PPTX パスまたはファイルライクオブジェクト（slides を渡す場合は不要）
        api_key: Anthropic API キー
        slides: 抽出済みスライド（extract_slides の返値）。テスト・再利用向け
        model: 不一致パスのモデル（既定 llm_models.MODEL_MISMATCH）
        typo_model: 誤字脱字パスのモデル（既定 llm_models.MODEL_TYPO）
        call_llm: LLM 呼び出しの差し替え（テスト用。call_anthropic と同じキーワード引数）
        progress: progress(done, total) で進捗を受け取るコールバック（メインスレッドから呼ぶ）

    Returns:
        dict: slides(枚数) / consistency(list) / typos(list) / errors(list[str]) /
              boilerplate(list) / models(dict)。所見1件ごとにも model を刷る。
              所見0件でもどのモデルで回したかが残るよう、models は結果dict側にも置く。
    """
    if slides is None:
        slides = extract_slides(pptx_file)
    models = {"consistency": model, "typos": typo_model}
    if not slides:
        return {"slides": 0, "consistency": [], "typos": [], "errors": [],
                "boilerplate": [], "models": models}

    call = call_llm or call_anthropic
    boilerplate = detect_boilerplate(slides)
    slide_map = {s["slide_num"]: s for s in slides}
    total = len(slides)

    tasks = []
    for chunk in _chunks(slides, chunk_size):
        tasks.append({
            "pass": "consistency", "chunk": chunk, "tool": CONSISTENCY_TOOL,
            "system": CONSISTENCY_SYSTEM, "model": model, "max_tokens": MAX_TOKENS_CONSISTENCY,
            "task_line": ("各スライドについて、まず原稿の要点を書き出し、"
                          "そのうえでスライド本文と食い違う点だけを report_consistency で報告してください。"),
        })
        tasks.append({
            "pass": "typos", "chunk": chunk, "tool": TYPO_TOOL,
            "system": TYPO_SYSTEM, "model": typo_model, "max_tokens": MAX_TOKENS_TYPO,
            "task_line": ("上記のスライド本文と原稿から、誤字・脱字・誤変換の可能性がある箇所だけを"
                          "report_typos で報告してください。"),
        })

    raw = {"consistency": [], "typos": []}
    errors = []
    done = 0
    if progress:
        progress(0, len(tasks))
    with ThreadPoolExecutor(max_workers=min(max_workers, len(tasks))) as pool:
        futures = {
            pool.submit(
                call,
                api_key=api_key, model=t["model"], system=t["system"], tool=t["tool"],
                max_tokens=t["max_tokens"],
                user_prompt=_chunk_prompt(t["chunk"], boilerplate, total, t["task_line"]),
            ): t
            for t in tasks
        }
        for future in as_completed(futures):
            task = futures[future]
            span = (f"S{task['chunk'][0]['slide_num']}〜S{task['chunk'][-1]['slide_num']}")
            label = "不一致" if task["pass"] == "consistency" else "誤字脱字"
            try:
                data = future.result()
                got = (
                    _flatten_consistency(data) if task["pass"] == "consistency"
                    else [f for f in (data or {}).get("findings") or [] if isinstance(f, dict)]
                )
                for f in got:
                    f["model"] = task["model"]
                raw[task["pass"]].extend(got)
            except Exception as e:
                # 例外クラス名を残す（2026-09-17 inspector指摘）: API起因の失敗も
                # 実装バグ（NameError等）も同じ文言に紛れると切り分けが遅れるため。
                errors.append(
                    f"{span} の{label}チェック（{task['model']}）が失敗しました: "
                    f"{type(e).__name__}: {e}"
                )
            done += 1
            if progress:
                progress(done, len(tasks))

    return {
        "slides": total,
        "consistency": _sanitize_findings(raw["consistency"], slide_map=slide_map,
                                          allowed_kinds=KINDS_CONSISTENCY),
        "typos": _sanitize_findings(raw["typos"], slide_map=slide_map,
                                    allowed_kinds=(KIND_TYPO,)),
        "errors": errors,
        "boilerplate": boilerplate,
        "models": models,
    }
