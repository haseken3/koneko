"""ステップ括り・Word 出力の回帰テスト（合成データのみ・API不要・$0）。

ここで固定するのは、教員の手元に出るものの形:
  1. ステップ分割**前**（steps=None）は1括りで出す（分割を強制しない）
  2. 所見0件のステップでも見出しと「気づいた点はありませんでした」を出す
     （見出しごと消えると、先生は「見てもらえた枚」と「何も出なかった枚」を区別できない）
  3. ステップの範囲外のスライド番号を**落とさない**（境界を手で動かすと steps が全枚を覆わない）
  4. 画面と Word が同じ `group_findings_by_step` で括る（片方だけ直る二重実装を作らない）
  5. Word の全 run に東アジアフォント（w:eastAsia）が入る（日本語が欧文代替で出る罠の回避）
  6. 確度「低」は本編に出ず、末尾の参考章にだけ出る

実行: /usr/bin/python3 -m pytest test_slide_script_report.py -q （koneko ルートから）
"""

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import slide_script_check_ui as ui
import slide_script_report as rep


def _finding(slide_num, *, kind="不足", confidence="高", where="notes", title="見出し"):
    return {
        "slide_num": slide_num, "kind": kind, "where": where,
        "quote": f"スライド{slide_num}の原稿にある要点です。",
        "counterpart": "", "reason": "スライド本文に見当たりません。",
        "suggestion": "1行足すと揃います。", "confidence": confidence,
        "quote_verified": True, "slide_title": title, "model": "test-model",
    }


def _step(num, label, start, end):
    return {"step_num": num, "step_label": label, "slide_range": (start, end),
            "first_title": "", "slide_count": end - start + 1, "char_count": 0,
            "pause_seconds": 0, "estimated_seconds": 0, "status": "ok"}


STEPS = [_step(1, "導入・学習目標", 1, 4), _step(2, "16進数の基礎", 5, 8)]


def _result(consistency=(), typos=()):
    return {"slides": 8, "consistency": list(consistency), "typos": list(typos),
            "errors": [], "boilerplate": [],
            "models": {"consistency": "model-a", "typos": "model-b"}}


def _docx_text(buf) -> str:
    return "\n".join(p.text for p in Document(buf).paragraphs)


# ─────────────────────────────────────────────
# 1. 分割前は1括り
# ─────────────────────────────────────────────
def test_without_steps_everything_is_one_group():
    groups = rep.group_findings_by_step([_finding(2), _finding(7)], None)
    assert len(groups) == 1
    assert groups[0]["heading"] == rep.ALL_SLIDES_HEADING
    assert "ステップ分割ができなかった" in groups[0]["heading"]
    assert [f["slide_num"] for f in groups[0]["findings"]] == [2, 7]


def test_without_steps_the_screen_shows_one_band():
    html = "".join(ui._group_html(g)
                   for g in rep.group_findings_by_step([_finding(2)], None))
    assert html.count("border-radius:6px;padding:10px 16px;margin:1.4rem") == 1
    assert rep.ALL_SLIDES_HEADING in html
    assert "気づいた点: 1件" in html
    assert "このステップで" not in html, "ステップ分割前なのに「このステップ」と言っている"


def test_without_steps_word_stays_one_group():
    text = _docx_text(rep.build_docx(_result(consistency=[_finding(2)]), None, "資料.pptx"))
    assert rep.ALL_SLIDES_HEADING in text
    assert "ステップ 1" not in text


# ─────────────────────────────────────────────
# 2. 0件のステップも見出しを出す
# ─────────────────────────────────────────────
def test_empty_step_keeps_its_heading_and_says_zero():
    groups = rep.group_findings_by_step([_finding(6)], STEPS)
    assert [g["heading"] for g in groups] == [
        "ステップ 1　導入・学習目標（スライド 1〜4）",
        "ステップ 2　16進数の基礎（スライド 5〜8）",
    ]
    assert groups[0]["findings"] == []

    html = ui._group_html(groups[0])
    assert "ステップ 1　導入・学習目標（スライド 1〜4）" in html
    assert rep.NO_FINDINGS_TEXT in html
    assert "このステップで気づいた点" not in html

    assert "このステップで気づいた点: 1件" in ui._group_html(groups[1])


def test_empty_step_keeps_its_heading_in_word():
    text = _docx_text(rep.build_docx(_result(consistency=[_finding(6)]), STEPS, "資料.pptx"))
    assert "ステップ 1　導入・学習目標（スライド 1〜4）" in text
    assert rep.NO_FINDINGS_TEXT in text


# ─────────────────────────────────────────────
# 3. 範囲外を落とさない
# ─────────────────────────────────────────────
def test_findings_outside_the_step_ranges_are_not_dropped():
    findings = [_finding(3), _finding(9), _finding(99)]
    groups = rep.group_findings_by_step(findings, STEPS)
    assert groups[-1]["heading"] == rep.OUTSIDE_STEPS_HEADING
    assert [f["slide_num"] for f in groups[-1]["findings"]] == [9, 99]

    kept = [f["slide_num"] for g in groups for f in g["findings"]]
    assert sorted(kept) == [3, 9, 99], "括り直しで所見が消えている"


def test_outside_group_is_absent_when_every_finding_fits():
    groups = rep.group_findings_by_step([_finding(3), _finding(6)], STEPS)
    assert [g["heading"] for g in groups] == [g["heading"] for g in groups[:2]]
    assert all(g["heading"] != rep.OUTSIDE_STEPS_HEADING for g in groups)


# ─────────────────────────────────────────────
# 4. 見出しの形
# ─────────────────────────────────────────────
def test_step_heading_shape():
    assert rep.step_heading(_step(2, "16進数の基礎", 5, 8)) == "ステップ 2　16進数の基礎（スライド 5〜8）"
    assert rep.step_heading(_step(3, "STEP3", 9, 9)) == "ステップ 3（スライド 9）", \
        "連番フォールバックのラベルが「ステップ 3　STEP3」と二重に出ている"


def test_slide_numbers_are_written_out_for_readers():
    """教員向けの面では「S5」の略記を使わない。"""
    card = ui._card_html(_finding(5))
    assert "スライド 5" in card
    assert ">S5<" not in card


# ─────────────────────────────────────────────
# 5. Word の体裁
# ─────────────────────────────────────────────
def test_every_word_run_declares_an_east_asian_font():
    doc = Document(rep.build_docx(_result(consistency=[_finding(6)]), STEPS, "資料.pptx"))
    runs = [r for p in doc.paragraphs for r in p.runs]
    assert runs, "本文が空"
    missing = [r.text for r in runs
               if r._element.rPr is None or r._element.rPr.rFonts is None
               or r._element.rPr.rFonts.get(qn("w:eastAsia")) != rep.EASTASIA_FONT]
    assert missing == [], f"東アジアフォント未指定の run がある: {missing[:3]}"


def test_word_carries_the_cover_and_the_disclaimer():
    text = _docx_text(rep.build_docx(_result(consistency=[_finding(6)]), STEPS, "第6回.pptx"))
    assert "第6回.pptx" in text
    assert rep.DISCLAIMER in text
    assert "最終的な判断は先生" in rep.DISCLAIMER, "断定しない約束が注意書きから消えている"
    assert "model-a" in text and "model-b" in text
    assert "□ 対応した" in text


def test_download_filename_keeps_the_material_name():
    from datetime import datetime
    name = rep.download_filename("第6回_プライバシー.pptx", now=datetime(2026, 9, 17, 15, 40))
    assert name == "第6回_プライバシー_チェック結果_20260917_1540.docx"


# ─────────────────────────────────────────────
# 6. 確度「低」の行き先
# ─────────────────────────────────────────────
def test_low_confidence_goes_to_the_reference_chapter_only():
    result = _result(consistency=[_finding(6), _finding(7, confidence="低")])
    text = _docx_text(rep.build_docx(result, STEPS, "資料.pptx"))
    assert "参考: 確度が低い気づき" in text
    body, reference = text.split("参考: 確度が低い気づき")
    assert "スライド 7" in reference
    assert "スライド 7" not in body, "確度が低い気づきが本編に出ている"


def test_reference_chapter_is_absent_when_nothing_is_low():
    text = _docx_text(rep.build_docx(_result(consistency=[_finding(6)]), STEPS, "資料.pptx"))
    assert "参考: 確度が低い気づき" not in text


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
