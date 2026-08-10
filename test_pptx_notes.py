"""ノート本文枠を失ったスライドの回帰テスト（2026-08-11・NeKo側で実測した型）。

実事象: 教員が PowerPoint でノート枠を消したスライドは `has_notes_slide` が True を返すのに
`notes_text_frame` が None になる（実測: 品質管理2_2ブロック260810・全43枚の5枚目。
教員提出73件2124スライド中1枚）。旧コードはそこで AttributeError で落ちていた。

その枚のノート面に残っていたのは「スライド イメージ」（空文字）と「スライド番号」（"5"）だけ
＝原稿は元から無いので、数える側は通してよい。
⚠ ただし素朴に「ノート面の全テキストを連結する」と **"5" が原稿として数えられる**。

ここで見るのは4点:
  1. 本文枠なしの枚を「ナレーション無し」として通す（尺推定が落ちない）
  2. スライド番号 "5" を原稿として数えない
  3. 本文枠は無いが文章が残っている個体は止める（黙って数え落とすと、尺が過小になり
     IDチェッカーは「ナレーションなし」という事実と違う所見を教員に返す）
  4. 読み手2箇所（narration_counter / idcheck.pptx_reader）が同じ判断に載っている

⚠ フィクスチャが実事象を再現できていること自体を毎回確かめる（修正前の書き方が
AttributeError で落ちることを固定する）＝再現しない土台の上の緑は意味がない。

実行: /usr/bin/python3 test_pptx_notes.py （koneko ルートから）／pytest でも走る
"""

import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
for _d in (str(ROOT), str(ROOT / "idcheck")):
    if _d not in sys.path:
        sys.path.insert(0, _d)

import narration_counter as nc
from pptx_notes import NotesBodyOrphanError, read_notes
from pptx_reader import read_slides

_NARRATION = "ここは先生が書いた原稿です。品質管理の考え方を説明します。"


def _make_pptx(*, orphan_text: str = "", slide_number_text: str = "5") -> str:
    """2枚のPPTXを作る。1枚目は正常、2枚目はノート本文枠を削った実事象の個体。

    orphan_text を渡すと、本文枠として認識されないまま文章が残っている個体になる。
    """
    prs = Presentation()
    blank = prs.slide_layouts[6]

    normal = prs.slides.add_slide(blank)
    normal.notes_slide.notes_text_frame.text = _NARRATION

    broken = prs.slides.add_slide(blank)
    notes = broken.notes_slide
    body = notes.notes_placeholder
    if orphan_text:
        body.text_frame.text = orphan_text
        body._element.find(".//" + qn("p:ph")).set("type", "obj")
    else:
        body._element.getparent().remove(body._element)  # ノート枠を消した状態
    for shape in notes.shapes:
        if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            shape.text_frame.text = slide_number_text

    path = Path(tempfile.mkdtemp()) / "notes_frame_missing.pptx"
    prs.save(str(path))
    return str(path)


def test_fixture_reproduces_the_real_individual():
    """フィクスチャが実事象と同じ形（has_notes_slide=True / notes_text_frame=None）になっている。"""
    slide = Presentation(_make_pptx()).slides[1]
    assert slide.has_notes_slide is True
    assert slide.notes_slide.notes_text_frame is None
    try:
        slide.notes_slide.notes_text_frame.text  # 修正前の書き方＝ここで落ちていた
    except AttributeError:
        pass
    else:
        raise AssertionError("修正前の書き方が落ちない＝実事象を再現できていない")
    print("✓ フィクスチャが実事象を再現している")


def test_narration_counter_passes_the_slide():
    """尺推定: 本文枠なしの枚で落ちず、ナレーション無しとして数える。"""
    result = nc.analyze_narration(_make_pptx())
    assert result["total_slides"] == 2
    assert result["slides_with_notes"] == 1
    assert result["slides"][1]["notes"] == ""
    assert result["slides"][1]["char_count"] == 0
    assert result["total_chars"] == nc._count_reading_chars(_NARRATION)
    print("✓ narration_counter: 本文枠なしを通す（尺は1枚目のみ）")


def test_slide_number_is_not_counted_as_narration():
    """ノート面に残る「スライド番号」を原稿として数えない（実測の落とし穴）。"""
    result = nc.analyze_narration(_make_pptx(slide_number_text="5"))
    assert result["slides"][1]["char_count"] == 0, "スライド番号を1文字数えている"
    assert result["slides_with_notes"] == 1
    print("✓ スライド番号を原稿として数えない")


def test_idcheck_reader_passes_the_slide():
    """IDチェッカー: 同じ個体で落ちず、notes が空になる。"""
    slides = read_slides(_make_pptx())
    assert len(slides) == 2
    assert slides[0]["notes"] == _NARRATION
    assert slides[1]["notes"] == ""
    print("✓ idcheck.pptx_reader: 本文枠なしを通す")


def test_orphan_narration_is_not_silently_dropped():
    """本文枠は無いが文章が残っている個体は止める（黙って数え落とさない）。"""
    path = _make_pptx(orphan_text="実は原稿がここに残っています。")

    for label, call in (("narration_counter", lambda: nc.analyze_narration(path)),
                        ("idcheck.pptx_reader", lambda: read_slides(path))):
        try:
            call()
        except NotesBodyOrphanError as e:
            assert "2枚中の2枚目" in str(e)                       # 何枚目かが人語で出る
            assert "実は原稿がここに残っています。" in str(e)
        else:
            raise AssertionError(f"{label}: 数え落としが素通りしている")
    print("✓ 数え落としになる個体は2経路とも止まる")


def test_both_readers_share_one_judgement():
    """読み手2箇所が同じ read_notes に載っている（別々に書くと片方だけ直る）。"""
    import pptx_reader
    assert nc.read_notes is read_notes
    assert pptx_reader.read_notes is read_notes
    print("✓ 読み手2箇所が同じ判断に載っている")


if __name__ == "__main__":
    test_fixture_reproduces_the_real_individual()
    test_narration_counter_passes_the_slide()
    test_slide_number_is_not_counted_as_narration()
    test_idcheck_reader_passes_the_slide()
    test_orphan_narration_is_not_silently_dropped()
    test_both_readers_share_one_judgement()
    print("\n✅ 全テスト通過（API不要・$0）")
